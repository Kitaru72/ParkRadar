from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent.parent / "output" / "presentation"
OUT.mkdir(parents=True, exist_ok=True)
TARGET = OUT / "ParkRadar_Hackathon_RU.pptx"

BG = RGBColor(244, 239, 231)
INK = RGBColor(25, 33, 38)
MUTED = RGBColor(94, 106, 115)
ACCENT = RGBColor(15, 118, 110)
ACCENT_DARK = RGBColor(11, 59, 54)
ACCENT_WARM = RGBColor(245, 158, 11)
CARD = RGBColor(255, 250, 240)
WHITE = RGBColor(255, 255, 255)


def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-0.7), Inches(-0.4), Inches(2.6), Inches(2.6))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.fill.transparency = 0.82
    shape.line.fill.background()

    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10.6), Inches(-0.2), Inches(2.2), Inches(2.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_WARM
    shape.fill.transparency = 0.78
    shape.line.fill.background()


def add_title(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(8.8), Inches(1.1))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = ACCENT_DARK
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.82), Inches(1.35), Inches(8.6), Inches(0.5))
        p2 = sub.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED


def add_card(slide, left, top, width, height, title, body, accent=ACCENT):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = accent
    card.line.transparency = 0.75

    tbox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.14), width - Inches(0.36), height - Inches(0.28))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = INK
    for line in body:
        para = tf.add_paragraph()
        para.text = line
        para.level = 0
        para.font.size = Pt(12)
        para.font.color.rgb = MUTED


def add_bullets(slide, left, top, width, height, lines):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20 if index == 0 else 18)
        p.font.bold = index == 0
        p.font.color.rgb = INK if index == 0 else MUTED


def add_kpi(slide, left, top, width, value, label):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, Inches(1.35))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.fill.background()
    tf = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.16), width - Inches(0.36), Inches(1.0)).text_frame
    p = tf.paragraphs[0]
    p.text = value
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = ACCENT_DARK
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.62), Inches(2.0), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = ACCENT_DARK
    pill.line.fill.background()
    p = slide.shapes.add_textbox(Inches(1.02), Inches(0.72), Inches(1.5), Inches(0.2)).text_frame.paragraphs[0]
    p.text = "Хакатон MVP"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = WHITE
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(6.4), Inches(2.0)).text_frame
    p = title.paragraphs[0]
    p.text = "ParkRadar"
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = ACCENT_DARK
    p = title.add_paragraph()
    p.text = "Понимаем парковку до того, как вы приедете"
    p.font.size = Pt(24)
    p.font.color.rgb = INK
    p = title.add_paragraph()
    p.text = "Сервис оценки свободных парковочных мест по адресу на основе публичных камер"
    p.font.size = Pt(16)
    p.font.color.rgb = MUTED
    add_card(slide, Inches(8.0), Inches(1.35), Inches(4.1), Inches(4.8), "Что показываем", [
        "рабочий пользовательский интерфейс",
        "admin UI и test-mode ingestion",
        "trip monitoring и scheduler",
        "строгое соответствие ТЗ",
    ])

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Проблема", "Парковка во дворе остается слепой зоной для жителя и управляющей стороны")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(3.8), Inches(4.6), "Житель", [
        "тратит время на круги по двору",
        "испытывает ежедневный стресс",
        "не понимает, есть ли смысл ехать к точке сейчас",
    ])
    add_card(slide, Inches(4.8), Inches(1.8), Inches(3.8), Inches(4.6), "УК / оператор", [
        "видит жалобы, но не видит цифры",
        "не понимает реальную загрузку зон",
        "не имеет data-driven основы для решений",
    ], ACCENT_WARM)
    add_card(slide, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.6), "Рынок", [
        "новое железо стоит дорого",
        "внедрение занимает время",
        "датчики не решают задачу быстрого поиска",
    ], ACCENT_DARK)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Решение", "Мы превращаем уже существующие публичные камеры в сервис понятного ответа")
    add_bullets(slide, Inches(0.9), Inches(1.8), Inches(5.8), Inches(4.8), [
        "Пользователь вводит адрес",
        "Система находит камеры в радиусе до 1 км",
        "Считает агрегированную доступность",
        "Показывает свободные места по классам машин",
        "Запускает trip monitoring на время поездки",
    ])
    add_card(slide, Inches(7.0), Inches(1.85), Inches(5.1), Inches(3.1), "Ключевой продуктовый принцип", [
        "не показываем видеопоток",
        "не заставляем пользователя анализировать камеры",
        "даем простой ответ для принятия решения",
    ])
    add_kpi(slide, Inches(7.0), Inches(5.2), Inches(1.55), "1 км", "радиус поиска")
    add_kpi(slide, Inches(8.75), Inches(5.2), Inches(1.55), "PWA", "web + Android")
    add_kpi(slide, Inches(10.5), Inches(5.2), Inches(1.55), "MVP", "строго по ТЗ")

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Что входит в MVP", "Реализованный контур соответствует границам ТЗ")
    items = [
        ("User flow", ["поиск по адресу или ссылке", "агрегированный ответ", "классы машин", "источники данных"]),
        ("Trip monitoring", ["создание trip_session", "half ETA", "pre-arrival", "fallback +15m"]),
        ("Admin UI", ["login по email/password", "реестр камер", "статистика", "ручное добавление камеры"]),
        ("Test ingestion", ["POST /test/screenshots", "POST /test/batch-screenshots", "ручной ввод free_*"]),
    ]
    x, y = Inches(0.8), Inches(1.8)
    for i, (title, body) in enumerate(items):
        add_card(slide, x + Inches((i % 2) * 6.0), y + Inches((i // 2) * 2.25), Inches(5.5), Inches(1.9), title, body)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Архитектура", "Demo-first архитектура, которая честно масштабируется в production")
    flow = [
        ("User PWA", 0.9, 2.1),
        ("Backend API", 4.5, 2.1),
        ("SQLite", 8.0, 2.1),
        ("Scheduler", 4.5, 4.4),
        ("Admin UI", 0.9, 4.4),
        ("Push / Pull", 8.0, 4.4),
    ]
    for name, x, y in flow:
        add_card(slide, Inches(x), Inches(y), Inches(2.2), Inches(1.3), name, [], ACCENT if "UI" in name or "PWA" in name else ACCENT_WARM)
    for x1, y1, x2, y2 in [(3.1, 2.75, 4.5, 2.75), (6.7, 2.75, 8.0, 2.75), (2.9, 5.0, 4.5, 5.0), (6.7, 5.0, 8.0, 5.0), (5.6, 3.4, 5.6, 4.4)]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = ACCENT_DARK
        line.line.width = Pt(2)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Демо-сценарий", "Как за 5-7 минут показать полный ценностный поток")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(2.8), Inches(3.8), "1. Пользователь", [
        "вводит адрес",
        "получает availability",
        "запускает trip monitoring",
    ])
    add_card(slide, Inches(3.9), Inches(1.8), Inches(2.8), Inches(3.8), "2. Админ", [
        "логинится",
        "обновляет snapshot",
        "запускает due-check",
    ], ACCENT_WARM)
    add_card(slide, Inches(7.0), Inches(1.8), Inches(2.8), Inches(3.8), "3. Система", [
        "пересчитывает availability",
        "создает уведомление",
        "обновляет trip session",
    ])
    add_card(slide, Inches(10.1), Inches(1.8), Inches(2.1), Inches(3.8), "4. Финал", [
        "пользователь видит изменение",
        "жюри видит end-to-end продукт",
    ], ACCENT_DARK)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Почему это сильное решение", "Не просто CV-демо, а продуктовый контур с реальным пользовательским сценарием")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.5), "Польза людям", [
        "меньше времени на поиск парковки",
        "меньше стресса и конфликтов",
        "больше предсказуемости поездки",
    ])
    add_card(slide, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.5), "Польза операторам", [
        "картина загрузки зон",
        "статистика запросов и hot spots",
        "основа для решений без нового железа",
    ], ACCENT_WARM)
    add_card(slide, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.5), "Польза для внедрения", [
        "быстрый MVP",
        "честный путь к pilot production",
        "модульная архитектура для scale",
    ], ACCENT_DARK)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Риски и честные ответы", "Мы не скрываем ограничения и сразу показываем mitigation")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(5.5), Inches(1.6), "Нестабильные камеры", ["status камеры", "degradation на доступный subset", "кэш валидных данных"])
    add_card(slide, Inches(6.0), Inches(1.8), Inches(6.0), Inches(1.6), "Точность в сложных условиях", ["в MVP test mode", "следующий шаг: production CV/ML", "дообучение под конкретный двор"], ACCENT_WARM)
    add_card(slide, Inches(0.8), Inches(3.7), Inches(5.5), Inches(1.6), "Юридическая рамка", ["используем только разрешенные публичные источники", "не показываем первичные кадры пользователю"])
    add_card(slide, Inches(6.0), Inches(3.7), Inches(6.0), Inches(1.6), "Надежность поездки", ["trip monitoring снижает ложные ожидания", "push architecture + pull fallback"], ACCENT_DARK)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Roadmap", "От hackathon MVP к pilot production и scale")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.6), "Этап 1. Сейчас", [
        "test-mode ingestion",
        "user PWA и admin UI",
        "trip monitoring",
        "встроенный scheduler",
    ])
    add_card(slide, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.6), "Этап 2. Pilot", [
        "реальные потоки камер",
        "точный CV/ML pipeline",
        "ETA от картографического API",
        "production push stack",
    ], ACCENT_WARM)
    add_card(slide, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.6), "Этап 3. Scale", [
        "worker для scheduler",
        "очереди задач",
        "Postgres и аналитика",
        "SLA / SLO / алерты",
    ], ACCENT_DARK)

    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Финал", "ParkRadar делает парковку предсказуемой, а не хаотичной")
    center = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.8), Inches(2.0))
    tf = center.text_frame
    p = tf.paragraphs[0]
    p.text = "Мы берем уже существующую визуальную инфраструктуру\nи превращаем ее в сервис принятия решения для человека."
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = ACCENT_DARK
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "Спасибо. Готовы показать живое демо."
    p.font.size = Pt(18)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER

    prs.save(TARGET)
    return TARGET


if __name__ == "__main__":
    path = build()
    print(path)
